"""Build a Word guide from a guidegen folder (steps.json + PNGs written by guide.py).

usage: tools\\guidegen\\.venv-docx\\Scripts\\python -I build_docx.py <guide_dir> <out.docx>

steps.json fields: title, subtitle, intro, before[], steps{id: {text, image?}}, verify{title,
text, code}, undo[]. Inline markup in any text: **bold**, `code`.
"""
from __future__ import annotations

import json
import re
import sys
import time
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x00, 0x3C, 0x77)  # Lab Console theme --brand-navy
GOLD = "F2B300"  # accent rule only, never text
GREY = RGBColor(0x59, 0x59, 0x59)
FONT, MONO = "Calibri", "Consolas"
IMG_W = Cm(15.5)

src, out = Path(sys.argv[1]), Path(sys.argv[2])
data = json.loads((src / "steps.json").read_text(encoding="utf-8"))
doc = Document()
for s in doc.sections:  # A4, 2 cm margins
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = FONT, Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 12)):
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.color.rgb, st.font.bold = FONT, Pt(size), NAVY, True


def runs(par, text: str, size: float | None = None) -> None:
    """Add text with **bold** and `code` markup."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = MONO
        else:
            r = par.add_run(part)
        if size:
            r.font.size = Pt(size)


def border(par, side: str, color: str, size: int = 12) -> None:
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", "4"), ("color", color)):
        el.set(qn(f"w:{k}"), v)
    bdr.append(el)
    ppr.append(bdr)


def shade(par, fill: str) -> None:
    shd = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", fill)):
        shd.set(qn(f"w:{k}"), v)
    par._p.get_or_add_pPr().append(shd)


def picture(path: Path, width: Cm = IMG_W) -> None:
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(14)
    inline = par.add_run().add_picture(str(path), width=width)
    sppr = inline._inline.graphic.graphicData.pic.spPr  # thin grey frame around the screenshot
    ln = OxmlElement("a:ln")
    ln.set("w", "9525")
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", "A6A6A6")
    fill.append(clr)
    ln.append(fill)
    sppr.append(ln)


def bullets(items: list[str], style: str = "List Bullet") -> None:
    for it in items:
        p = doc.add_paragraph(style=style)
        runs(p, it)


def code(text: str) -> None:
    lines = text.split("\n")
    for i, line in enumerate(lines):  # one paragraph per line (never \n inside a run)
        p = doc.add_paragraph()
        shade(p, "F2F2F2")
        pf = p.paragraph_format
        pf.space_after = pf.space_before = Pt(0)
        pf.keep_with_next = i < len(lines) - 1
        r = p.add_run(line or " ")
        r.font.name, r.font.size = MONO, Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(rows: list[list[str]], widths: list[float] | None = None, mono: list[int] = ()) -> None:
    """First row = header (navy, white bold). widths in cm. mono: column indexes shown verbatim in
    a monospace font (no markup), e.g. Splunk searches whose macros use backticks."""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            if widths:
                cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i > 0 and j in mono:
                r = p.add_run(val)
                r.font.name, r.font.size = MONO, Pt(8.5)
            else:
                runs(p, val, size=9.5)
            if i == 0:
                tcpr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                for k, v in (("val", "clear"), ("color", "auto"), ("fill", "003C77")):
                    shd.set(qn(f"w:{k}"), v)
                tcpr.append(shd)
                for r in p.runs:
                    r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def sections(secs: list[dict]) -> None:
    """[{heading, items: [{p}|{bullets}|{numbered}|{table, widths?}|{image, width_cm?}|{code}|{h2}]}]"""
    for sec in secs:
        doc.add_paragraph(sec["heading"], style="Heading 1").paragraph_format.keep_with_next = True
        for it in sec["items"]:
            if "p" in it:
                runs(doc.add_paragraph(), it["p"])
            elif "h2" in it:
                doc.add_paragraph(it["h2"], style="Heading 2").paragraph_format.keep_with_next = True
            elif "bullets" in it:
                bullets(it["bullets"])
            elif "numbered" in it:
                bullets(it["numbered"], "List Number")
            elif "table" in it:
                table(it["table"], it.get("widths"), it.get("mono", []))
            elif "image" in it:
                picture(src / it["image"], Cm(it.get("width_cm", 15.5)))
            elif "code" in it:
                code(it["code"])


title = doc.add_paragraph(data.get("title", "Guide"), style="Title")
border(title, "bottom", GOLD, 18)
if data.get("subtitle"):
    p = doc.add_paragraph()
    r = p.add_run(data["subtitle"])
    r.font.color.rgb, r.font.size = GREY, Pt(10)
if data.get("intro"):
    runs(doc.add_paragraph(), data["intro"])
if data.get("before"):
    doc.add_paragraph("Before you start", style="Heading 1")
    bullets(data["before"])
sections(data.get("sections_before", []))

doc.add_paragraph(data.get("steps_heading", "Steps"), style="Heading 1")
for n, sid in enumerate(sorted(data["steps"]), 1):
    st = data["steps"][sid]
    if st.get("group"):
        doc.add_paragraph(st["group"], style="Heading 2").paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"Step {n}   ")
    r.bold, r.font.color.rgb = True, NAVY
    runs(p, st["text"])
    if st.get("image"):
        picture(src / st["image"])

sections(data.get("sections_after", []))
if data.get("verify"):
    v = data["verify"]
    h = doc.add_paragraph(v["title"], style="Heading 1")
    h.paragraph_format.keep_with_next = True
    if v.get("text"):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        runs(p, v["text"])
    lines = v["code"].split("\n")
    for i, line in enumerate(lines):  # one paragraph per line (never \n inside a run)
        p = doc.add_paragraph()
        shade(p, "F2F2F2")
        pf = p.paragraph_format
        pf.space_after = pf.space_before = Pt(0)
        pf.keep_with_next = i < len(lines) - 1
        r = p.add_run(line or " ")
        r.font.name, r.font.size = MONO, Pt(9)
        if i == 0:
            r.bold = True
if data.get("undo"):
    doc.add_paragraph("Undo", style="Heading 1")
    bullets(data["undo"])

foot = doc.sections[0].footer.paragraphs[0]
fr = foot.add_run(f"{data.get('title', '')}  |  {time.strftime('%Y-%m-%d')}")
fr.font.size, fr.font.color.rgb = Pt(8), GREY
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(out))
print(f"wrote {out} ({out.stat().st_size} bytes, {len(data['steps'])} steps)")
